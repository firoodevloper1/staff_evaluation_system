from decimal import Decimal
from io import BytesIO

from django.contrib import messages
from django.core.exceptions import ValidationError
from django.contrib.auth import login, logout
from django.contrib.auth.decorators import login_required
from django.contrib.auth.forms import AuthenticationForm
from django.contrib.auth.views import LoginView
from django.db.models import Q
from django.http import HttpResponse, HttpResponseBadRequest
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone
from django.utils.dateparse import parse_date
from django.views.decorators.http import require_POST
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from .criteria import FORM_TITLES, RATING_OPTIONS, WEIGHTS, get_catalog
from .forms import SignUpForm, StaffMemberForm, AdministrativeEvaluationForm, FinalReviewForm, ArchiveEntryForm, EvaluationRecordForm
from .models import EvaluationRecord, AdministrativeEvaluation, StaffMember, UserProfile, FinalEvaluationReview, ArchiveEntry
from .scoring import average_rating, category_average

MANAGEMENT_ROLES = {"department_head", "committee", "college_admin", "dean", "academic_vp", "archivist", "office_directorate"}
ADMINISTRATIVE_ROLES = {"department_head", "committee", "college_admin"}
EXPORT_ROLES = MANAGEMENT_ROLES


def get_role(user):
    if not user.is_authenticated:
        return ""
    if user.is_superuser:
        return "admin"
    try:
        return user.evaluation_profile.role
    except Exception:
        return "instructor"


def is_manager(user):
    return user.is_authenticated and (user.is_superuser or get_role(user) in MANAGEMENT_ROLES)


def is_global_manager(user):
    return user.is_authenticated and (user.is_superuser or get_role(user) in {"college_admin", "dean", "academic_vp", "archivist", "office_directorate"})


def scoped_staff(user, include_inactive=False):
    qs = StaffMember.objects.all() if include_inactive else StaffMember.objects.filter(is_active=True)
    role = get_role(user)
    if user.is_superuser or role in {"college_admin", "dean", "academic_vp", "archivist", "office_directorate"}:
        return qs
    if role in {"department_head", "committee"}:
        try:
            department = user.evaluation_profile.department.strip()
            college = user.evaluation_profile.college.strip()
        except Exception:
            department, college = "", ""
        if department:
            qs = qs.filter(department__iexact=department)
        elif college and college != "Ambo University":
            qs = qs.filter(college__iexact=college)
        else:
            return qs.none()
        return qs
    return StaffMember.objects.filter(linked_user=user) if user.is_authenticated else qs.none()


def is_administrative_manager(user):
    return user.is_authenticated and (user.is_superuser or get_role(user) in ADMINISTRATIVE_ROLES)


def can_archive(user):
    archive_roles = {"department_head", "committee", "college_admin", "dean", "academic_vp", "archivist"}
    return user.is_authenticated and (user.is_superuser or get_role(user) in archive_roles)


def visible_archive_entries(user):
    qs = ArchiveEntry.objects.select_related(
        "academic_record__staff", "administrative_record__staff", "archived_by"
    )
    if not is_manager(user):
        return qs.none()
    if is_global_manager(user):
        return qs
    if get_role(user) in {"department_head", "committee"}:
        staff_scope = scoped_staff(user, include_inactive=True)
        return qs.filter(
            Q(academic_record__staff__in=staff_scope)
            | Q(administrative_record__staff__in=staff_scope)
        ).distinct()
    return qs.none()


def can_create_form(user, form_type):
    role = get_role(user)
    allowed = {
        "student": {"student"},
        "department_head": {"department_head", "committee"},
        "peer": {"peer", "department_head", "committee"},
        "self": {"instructor"},
    }
    return user.is_superuser or role in allowed.get(form_type, set())


def visible_evaluations(user):
    qs = EvaluationRecord.objects.select_related("staff", "evaluator")
    if is_global_manager(user):
        return qs
    if get_role(user) in {"department_head", "committee"}:
        staff_scope = scoped_staff(user, include_inactive=True)
        return qs.filter(staff__in=staff_scope)
    return qs.filter(evaluator=user)


def visible_administrative_evaluations(user):
    qs = AdministrativeEvaluation.objects.select_related("staff", "evaluator")
    if is_global_manager(user):
        return qs
    if get_role(user) in {"department_head", "committee"}:
        staff_scope = scoped_staff(user, include_inactive=True)
        return qs.filter(staff__in=staff_scope)
    return qs.filter(evaluator=user)


def render_forbidden(request):
    return render(request, "403.html", status=403)


def staff_directory_data(queryset=None):
    queryset = queryset if queryset is not None else StaffMember.objects.filter(is_active=True)
    return {str(item["id"]): item for item in queryset.values("id", "position", "grade", "job_process", "department", "college")}


def home(request):
    if request.user.is_authenticated:
        return redirect("dashboard")
    return render(request, "home.html")


class SignInView(LoginView):
    template_name = "registration/login.html"
    authentication_form = AuthenticationForm
    redirect_authenticated_user = True


def signup(request):
    if request.user.is_authenticated:
        return redirect("dashboard")
    form = SignUpForm(request.POST or None)
    if request.method == "POST" and form.is_valid():
        user = form.save()
        login(request, user)
        messages.success(request, "Your account has been created. An administrator can assign additional permissions if needed.")
        return redirect("dashboard")
    return render(request, "registration/signup.html", {"form": form})


@require_POST
def signout(request):
    logout(request)
    return redirect("home")


@login_required
def dashboard(request):
    role = get_role(request.user)
    records = visible_evaluations(request.user).filter(status="submitted")
    admin_records = visible_administrative_evaluations(request.user).filter(status="submitted")
    staff_qs = scoped_staff(request.user) if is_manager(request.user) else StaffMember.objects.filter(linked_user=request.user, is_active=True)
    if not is_manager(request.user):
        records = records.filter(evaluator=request.user)
        admin_records = admin_records.filter(evaluator=request.user)
    form_counts = {key: records.filter(form_type=key).count() for key in WEIGHTS}
    recent = records.order_by("-created_at")[:6]
    total = records.count() + admin_records.count()
    staff_count = staff_qs.count() if is_manager(request.user) else StaffMember.objects.filter(linked_user=request.user).count()
    avg_values = [float(r.average_score) for r in records[:500]]
    average_score = round(sum(avg_values) / len(avg_values), 2) if avg_values else 0
    chart_data = [
        {"label": label, "count": form_counts[key], "color": color}
        for key, label, color in [
            ("student", "Students", "blue"),
            ("department_head", "Head", "amber"),
            ("peer", "Peers", "purple"),
            ("self", "Self", "green"),
        ]
    ]
    context = {
        "role": role, "total_records": total, "staff_count": staff_count, "today": timezone.localdate(),
        "academic_records": records.count(), "administrative_records": admin_records.count(),
        "average_score": average_score, "form_counts": form_counts,
        "recent_records": recent, "chart_data": chart_data,
        "can_administer": is_administrative_manager(request.user),
        "can_manage": is_manager(request.user),
    }
    return render(request, "dashboard.html", context)


@login_required
def evaluation_create(request, form_type):
    if form_type not in WEIGHTS:
        return HttpResponseBadRequest("Unknown evaluation form type.")
    if not can_create_form(request.user, form_type):
        return render_forbidden(request)
    catalog = get_catalog(form_type)
    setup_message = ""
    if get_role(request.user) in {"department_head", "committee"} and not request.user.is_superuser:
        staff_members = scoped_staff(request.user).order_by("name")
    else:
        staff_members = StaffMember.objects.filter(is_active=True).order_by("name")
    if form_type == "self":
        own = StaffMember.objects.filter(linked_user=request.user, is_active=True).first()
        if own:
            staff_members = StaffMember.objects.filter(pk=own.pk)
        else:
            staff_members = StaffMember.objects.none()
            setup_message = "Your account is not linked to a staff directory record yet. Ask the administrator to link your account before submitting a self-evaluation."
    selected_ratings = {}
    error_message = ""

    if request.method == "POST":
        staff_id = request.POST.get("staff_id")
        staff = staff_members.filter(pk=staff_id).first()
        if not staff:
            error_message = "Select a valid staff member. If this is a self-evaluation, ask the administrator to link your staff profile to your account."
        elif form_type == "self" and staff.linked_user_id and staff.linked_user_id != request.user.id:
            error_message = "Self-evaluation must be submitted for your own linked staff profile."
        else:
            submitted_ratings = {}
            missing = []
            for category in catalog:
                for item in category["items"]:
                    name = "rating_" + item["key"]
                    selected = request.POST.get(name, "").strip()
                    selected_ratings[item["key"]] = selected
                    if selected not in {"5", "4", "3", "2", "1", "0", "NA"}:
                        missing.append(item["key"])
                    else:
                        submitted_ratings[item["key"]] = selected
            if missing:
                error_message = f"Please answer all rating items before submitting. {len(missing)} item(s) remain unanswered."
            else:
                eval_date = parse_date(request.POST.get("evaluation_date", "")) or timezone.localdate()
                period_start = parse_date(request.POST.get("period_start", ""))
                period_end = parse_date(request.POST.get("period_end", ""))
                academic_year = request.POST.get("academic_year", "").strip()
                if period_start and period_end and period_end < period_start:
                    error_message = "The evaluation end date must not be earlier than its start date."
                else:
                    record = EvaluationRecord.objects.create(
                        form_type=form_type,
                        staff=staff,
                        evaluator=request.user,
                        college=request.POST.get("college", staff.college).strip() or staff.college,
                        department=request.POST.get("department", staff.department).strip() or staff.department,
                        course_title=request.POST.get("course_title", "").strip(),
                        course_code=request.POST.get("course_code", "").strip(),
                        academic_year=academic_year,
                        batch=request.POST.get("batch", ""),
                        semester=request.POST.get("semester", ""),
                        evaluation_date=eval_date,
                        period_start=period_start,
                        period_end=period_end,
                        ratings=submitted_ratings,
                        average_score=average_rating(submitted_ratings),
                        strengths=request.POST.get("strengths", "").strip(),
                        improvements=request.POST.get("improvements", "").strip(),
                        status="submitted",
                        submitted_at=timezone.now(),
                    )
                    messages.success(request, "Evaluation submitted successfully.")
                    return redirect("evaluation_detail", pk=record.pk)

    for category in catalog:
        for item in category["items"]:
            item["selected"] = selected_ratings.get(item["key"], "")
    return render(request, "evaluation_form.html", {
        "form_type": form_type,
        "form_title": FORM_TITLES[form_type],
        "weight": WEIGHTS[form_type],
        "catalog": catalog,
        "staff_members": staff_members,
        "rating_options": RATING_OPTIONS,
        "error_message": error_message,
        "setup_message": setup_message,
        "selected_ratings": selected_ratings,
        "today": timezone.localdate().isoformat(),
        "form_metadata": form_type in {"student", "department_head", "peer", "self"},
        "selected_staff_id": request.POST.get("staff_id", request.GET.get("staff", "")),
        "staff_directory_data": staff_directory_data(staff_members),
        "post_values": request.POST if request.method == "POST" else {},
    })


@login_required
def evaluation_list(request):
    records = visible_evaluations(request.user)
    if form_type := request.GET.get("form_type", "").strip():
        if form_type in WEIGHTS:
            records = records.filter(form_type=form_type)
    if status := request.GET.get("status", "").strip():
        if status in {"draft", "submitted"}:
            records = records.filter(status=status)
    if year := request.GET.get("year", "").strip():
        records = records.filter(academic_year__icontains=year)
    if query := request.GET.get("q", "").strip():
        records = records.filter(
            Q(staff__name__icontains=query) | Q(staff__employee_id__icontains=query)
            | Q(course_title__icontains=query) | Q(course_code__icontains=query)
            | Q(department__icontains=query) | Q(academic_year__icontains=query)
        )
    records = records.order_by("-created_at")
    return render(request, "evaluation_list.html", {
        "records": records[:400], "form_type_filter": request.GET.get("form_type", ""),
        "status_filter": request.GET.get("status", ""), "year_filter": request.GET.get("year", ""),
        "query": request.GET.get("q", ""), "can_export": is_manager(request.user),
    })


@login_required
def evaluation_detail(request, pk):
    record = get_object_or_404(visible_evaluations(request.user), pk=pk)
    catalog = get_catalog(record.form_type)
    for category in catalog:
        category["average"] = category_average(category, record.ratings)
    return render(request, "evaluation_detail.html", {
        "record": record, "catalog": catalog, "form_title": FORM_TITLES[record.form_type],
        "weight": record.weight, "weighted_score": record.weighted_score,
        "can_archive": can_archive(request.user),
    })


@login_required
def evaluation_update(request, pk):
    record = get_object_or_404(visible_evaluations(request.user), pk=pk)
    catalog = get_catalog(record.form_type)
    for category in catalog:
        for item in category["items"]:
            item["selected"] = record.ratings.get(item["key"], "") if record.ratings else ""
    selected_ratings = {item["key"]: item["selected"] for category in catalog for item in category["items"]}

    error_message = ""
    if request.method == "POST":
        staff_id = request.POST.get("staff_id")
        staff = StaffMember.objects.filter(pk=staff_id).first() if staff_id else record.staff
        if not staff:
            error_message = "Select a valid staff member."
        else:
            submitted_ratings = {}
            missing = []
            for category in catalog:
                for item in category["items"]:
                    name = "rating_" + item["key"]
                    selected = request.POST.get(name, "").strip()
                    selected_ratings[item["key"]] = selected
                    if selected not in {"5", "4", "3", "2", "1", "0", "NA"}:
                        missing.append(item["key"])
                    else:
                        submitted_ratings[item["key"]] = selected
            if missing and not request.POST.get("strengths") and not request.POST.get("improvements"):
                error_message = f"Please answer all rating items before submitting. {len(missing)} item(s) remain unanswered."
            else:
                eval_date = parse_date(request.POST.get("evaluation_date", "")) or timezone.localdate()
                period_start = parse_date(request.POST.get("period_start", ""))
                period_end = parse_date(request.POST.get("period_end", ""))
                academic_year = request.POST.get("academic_year", "").strip()
                if period_start and period_end and period_end < period_start:
                    error_message = "The evaluation end date must not be earlier than its start date."
                else:
                    record.staff = staff
                    record.evaluator = request.user
                    record.college = request.POST.get("college", staff.college).strip() or staff.college
                    record.department = request.POST.get("department", staff.department).strip() or staff.department
                    record.course_title = request.POST.get("course_title", "").strip()
                    record.course_code = request.POST.get("course_code", "").strip()
                    record.academic_year = academic_year
                    record.batch = request.POST.get("batch", "")
                    record.semester = request.POST.get("semester", "")
                    record.evaluation_date = eval_date
                    record.period_start = period_start
                    record.period_end = period_end
                    record.ratings = submitted_ratings if submitted_ratings else record.ratings
                    record.average_score = average_rating(submitted_ratings) if submitted_ratings else record.average_score
                    record.strengths = request.POST.get("strengths", "").strip()
                    record.improvements = request.POST.get("improvements", "").strip()
                    record.status = request.POST.get("status", "submitted")
                    record.save()
                    messages.success(request, "Evaluation updated successfully.")
                    return redirect("evaluation_detail", pk=record.pk)

    for category in catalog:
        for item in category["items"]:
            item["selected"] = selected_ratings.get(item["key"], "")

    return render(request, "evaluation_form.html", {
        "form_type": record.form_type,
        "form_title": FORM_TITLES[record.form_type],
        "weight": record.weight,
        "catalog": catalog,
        "staff_members": StaffMember.objects.filter(is_active=True).order_by("name"),
        "rating_options": RATING_OPTIONS,
        "error_message": error_message,
        "setup_message": "",
        "selected_ratings": selected_ratings,
        "today": timezone.localdate().isoformat(),
        "form_metadata": record.form_type in {"student", "department_head", "peer", "self"},
        "selected_staff_id": record.staff.pk,
        "staff_directory_data": staff_directory_data(StaffMember.objects.filter(is_active=True)),
        "post_values": {},
    })


# --- Excel helpers and exports -------------------------------------------------
def styled_workbook(title, headers):
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append(headers)
    header_fill = PatternFill("solid", fgColor="0B427A")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="medium", color="F7BE68"))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 32
    return wb, ws


def finish_workbook_response(wb, filename):
    content = BytesIO()
    wb.save(content)
    content.seek(0)
    response = HttpResponse(
        content.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def format_sheet(ws, max_width=52):
    for col in ws.columns:
        letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            value = str(cell.value or "")
            max_len = max(max_len, min(len(value), max_width))
        ws.column_dimensions[letter].width = min(max(max_len + 2, 12), max_width)


@login_required
def export_evaluations(request):
    if not is_manager(request.user):
        return render_forbidden(request)
    records = visible_evaluations(request.user)
    form_type = request.GET.get("form_type", "").strip()
    if form_type in WEIGHTS:
        records = records.filter(form_type=form_type)
    if year := request.GET.get("year", "").strip():
        records = records.filter(academic_year__icontains=year)
    status = request.GET.get("status", "").strip()
    if status in {"draft", "submitted"}:
        records = records.filter(status=status)
    if query := request.GET.get("q", "").strip():
        records = records.filter(Q(staff__name__icontains=query) | Q(staff__employee_id__icontains=query) | Q(department__icontains=query))
    records = list(records.order_by("staff__name", "-evaluation_date"))
    selected_types = sorted({r.form_type for r in records})
    headers = ["Form type", "Employee ID", "Instructor name", "College", "Department", "Evaluator", "Course title", "Course code", "Academic year", "Batch", "Semester", "Evaluation date", "Period start", "Period end", "Average / 5", "Weighted contribution", "Strengths", "Improvement points", "Status"]
    item_columns = []
    for typ in selected_types:
        for category in get_catalog(typ):
            for item in category["items"]:
                item_columns.append((typ, item["key"], f"{FORM_TITLES[typ]} — {category['title']}: {item['text']}"))
    dynamic_headers = [label for _typ, _key, label in item_columns]
    wb, ws = styled_workbook("Evaluation data", headers + dynamic_headers)
    for r in records:
        row = [
            r.get_form_type_display(), r.staff.employee_id, r.staff.name, r.college, r.department,
            r.evaluator.get_full_name() or r.evaluator.username, r.course_title, r.course_code,
            r.academic_year, r.get_batch_display() if r.batch else "", r.get_semester_display() if r.semester else "",
            r.evaluation_date.isoformat(), r.period_start.isoformat() if r.period_start else "",
            r.period_end.isoformat() if r.period_end else "", float(r.average_score), float(r.weighted_score),
            r.strengths, r.improvements, r.get_status_display(),
        ]
        row.extend(r.ratings.get(key, "") if r.form_type == typ else "" for typ, key, _label in item_columns)
        ws.append(row)
    ws.auto_filter.ref = ws.dimensions
    format_sheet(ws)
    suffix = form_type if form_type in WEIGHTS else "all"
    return finish_workbook_response(wb, f"ambo_evaluations_{suffix}_{timezone.localdate().isoformat()}.xlsx")


@login_required
def export_administrative(request):
    if not is_manager(request.user):
        return render_forbidden(request)
    records = visible_administrative_evaluations(request.user)
    status = request.GET.get("status", "").strip()
    if status in {"draft", "submitted"}:
        records = records.filter(status=status)
    if query := request.GET.get("q", "").strip():
        records = records.filter(Q(staff__name__icontains=query) | Q(staff__employee_id__icontains=query) | Q(staff__department__icontains=query))
    headers = [
        "Employee ID", "Employee full name", "Job title", "Grade", "Work process", "College", "Department",
        "Period start", "Period end", "Evaluation plan", "Planned outcomes", "Actual achievements",
        "Plan execution / 70", "Self / 5", "Team / 15", "Supervisor / 10", "Total / 100", "Performance level",
        "Strengths", "Improvement gaps", "Supervisor name", "Supervisor signature (typed)", "Supervisor date",
        "Employee agreed", "Reason for disagreement", "Employee signature (typed)", "Employee date", "Evaluator", "Status",
    ]
    wb, ws = styled_workbook("Six-month admin", headers)
    for r in records.order_by("-period_end"):
        ws.append([
            r.staff.employee_id, r.staff.name, r.staff.position, r.staff.grade, r.staff.job_process,
            r.staff.college, r.staff.department, r.period_start.isoformat(), r.period_end.isoformat(),
            r.plan_title, r.planned_outcomes, r.actual_achievements, float(r.plan_execution_score),
            float(r.employee_score), float(r.team_score), float(r.supervisor_score), float(r.total_score),
            r.get_performance_level_display(), r.strengths, r.improvement_gaps, r.supervisor_name,
            r.supervisor_signature, r.supervisor_date.isoformat() if r.supervisor_date else "",
            "Yes" if r.employee_agreed else "No", r.disagreement_reason, r.employee_signature,
            r.employee_date.isoformat() if r.employee_date else "", r.evaluator.get_full_name() or r.evaluator.username,
            r.get_status_display(),
        ])
    format_sheet(ws)
    return finish_workbook_response(wb, f"ambo_six_month_evaluations_{timezone.localdate().isoformat()}.xlsx")


@login_required
def summary(request):
    staff_members = scoped_staff(request.user) if is_manager(request.user) else StaffMember.objects.filter(linked_user=request.user, is_active=True)
    staff_id = request.GET.get("staff", "").strip()
    selected_staff = staff_members.filter(pk=staff_id).first() if staff_id.isdigit() else None
    year = request.GET.get("year", "").strip()
    if not is_manager(request.user):
        linked = StaffMember.objects.filter(linked_user=request.user, is_active=True).first()
        if linked:
            staff_members = staff_members.filter(pk=linked.pk)
            selected_staff = linked
        else:
            staff_members = StaffMember.objects.none()
    elif selected_staff and not staff_members.filter(pk=selected_staff.pk).exists():
        selected_staff = None
    years = EvaluationRecord.objects.exclude(academic_year="").values_list("academic_year", flat=True).distinct().order_by("-academic_year")
    if selected_staff and get_role(request.user) == "instructor" and selected_staff.linked_user_id == request.user.id:
        # Instructors can see their consolidated summary while raw individual records stay restricted.
        records = EvaluationRecord.objects.filter(staff=selected_staff, status="submitted")
    else:
        records = visible_evaluations(request.user).filter(status="submitted")
    if selected_staff:
        records = records.filter(staff=selected_staff)
    if year:
        selected_year = year
        records = records.filter(academic_year=selected_year)
    elif selected_staff:
        selected_year = records.exclude(academic_year="").order_by("-academic_year").values_list("academic_year", flat=True).first() or str(timezone.localdate().year)
        records = records.filter(academic_year=selected_year)
    elif not is_manager(request.user):
        records = records.none()
        selected_year = str(timezone.localdate().year)
    else:
        selected_year = str(timezone.localdate().year)

    forms_summary = []
    total_weighted = Decimal("0.00")
    catalog_summary = []
    for form_type, weight in WEIGHTS.items():
        form_records = list(records.filter(form_type=form_type))
        if form_records:
            mean = (sum((r.average_score for r in form_records), Decimal("0")) / Decimal(len(form_records))).quantize(Decimal("0.01"))
            contribution = (mean * Decimal(weight) / Decimal("5")).quantize(Decimal("0.01"))
        else:
            mean, contribution = Decimal("0.00"), Decimal("0.00")
        total_weighted += contribution
        category_rows = []
        for category in get_catalog(form_type):
            scores = []
            for rec in form_records:
                value = category_average(category, rec.ratings)
                if value is not None:
                    scores.append(value)
            category_mean = (sum(scores, Decimal("0")) / Decimal(len(scores))).quantize(Decimal("0.01")) if scores else None
            category_rows.append({"title": category["title"], "average": category_mean})
        forms_summary.append({
            "form_type": form_type, "title": FORM_TITLES[form_type], "weight": weight,
            "count": len(form_records), "average": mean, "contribution": contribution,
            "categories": category_rows,
        })
    total_weighted = min(total_weighted, Decimal("100.00")).quantize(Decimal("0.01"))

    review = None
    review_form = FinalReviewForm()
    if selected_staff:
        review = FinalEvaluationReview.objects.filter(staff=selected_staff, academic_year=selected_year).first()
        if review:
            review_form = FinalReviewForm(instance=review)
    if request.method == "POST":
        if not is_administrative_manager(request.user) or not selected_staff:
            return render_forbidden(request)
        review_form = FinalReviewForm(request.POST, instance=review)
        if review_form.is_valid():
            saved = review_form.save(commit=False)
            saved.staff = selected_staff
            saved.academic_year = selected_year
            saved.reviewed_by = request.user
            saved.reviewed_at = timezone.now()
            saved.save()
            messages.success(request, "Final evaluation review saved.")
            return redirect(f"{request.path}?staff={selected_staff.pk}&year={selected_year}")
    return render(request, "summary.html", {
        "staff_members": staff_members, "selected_staff": selected_staff, "year": year,
        "years": years, "selected_year": selected_year, "forms_summary": forms_summary,
        "total_weighted": total_weighted, "review": review, "review_form": review_form,
        "can_review": is_administrative_manager(request.user) and selected_staff is not None,
        "can_manage": is_manager(request.user),
        "can_export_summary": selected_staff is not None and (is_manager(request.user) or (get_role(request.user) == "instructor" and selected_staff.linked_user_id == request.user.id)),
        "summary_only_categories": ["Performing non-instructional duties"],
    })


@login_required
def export_summary(request):
    staff_id = request.GET.get("staff", "").strip()
    year = request.GET.get("year", "").strip()
    if is_manager(request.user):
        staff_scope = scoped_staff(request.user, include_inactive=True)
        staff = staff_scope.filter(pk=staff_id).first() if staff_id.isdigit() else None
        if staff_id.isdigit() and staff is None:
            return render_forbidden(request)
        records = visible_evaluations(request.user).filter(status="submitted")
    elif get_role(request.user) == "instructor":
        staff = StaffMember.objects.filter(linked_user=request.user, is_active=True, pk=staff_id).first() if staff_id.isdigit() else None
        if staff is None:
            return render_forbidden(request)
        records = EvaluationRecord.objects.filter(staff=staff, status="submitted")
    else:
        return render_forbidden(request)
    if staff:
        records = records.filter(staff=staff)
    if year:
        records = records.filter(academic_year=year)
    headers = ["Instructor", "Employee ID", "Academic year", "Evaluator group", "Evaluation count", "Average / 5", "Group weight", "Weighted result", "Focus area", "Focus-area average / 5"]
    wb, ws = styled_workbook("Form E summary", headers)
    final_total = Decimal("0.00")
    groups = []
    for form_type, weight in WEIGHTS.items():
        form_records = list(records.filter(form_type=form_type))
        avg = (sum((r.average_score for r in form_records), Decimal("0")) / len(form_records)).quantize(Decimal("0.01")) if form_records else Decimal("0.00")
        contribution = (avg * Decimal(weight) / Decimal("5")).quantize(Decimal("0.01"))
        final_total += contribution
        focus_averages = []
        for category in get_catalog(form_type):
            scores = [value for rec in form_records if (value := category_average(category, rec.ratings)) is not None]
            cat_avg = (sum(scores, Decimal("0")) / len(scores)).quantize(Decimal("0.01")) if scores else ""
            focus_averages.append((category["title"], cat_avg))
        groups.append((form_type, form_records, avg, contribution, focus_averages))
    for form_type, form_records, avg, contribution, focus_averages in groups:
        for focus, cat_avg in focus_averages:
            ws.append([staff.name if staff else "All instructors", staff.employee_id if staff else "", year, FORM_TITLES[form_type], len(form_records), float(avg), WEIGHTS[form_type], float(contribution), focus, float(cat_avg) if cat_avg != "" else ""])
        if not focus_averages:
            ws.append([staff.name if staff else "All instructors", staff.employee_id if staff else "", year, FORM_TITLES[form_type], len(form_records), float(avg), WEIGHTS[form_type], float(contribution), "", ""])
    ws.append([])
    ws.append(["Total result out of 100", "", year, "", "", "", "", float(min(final_total, Decimal("100"))), "", ""])
    format_sheet(ws)
    safe_id = staff.employee_id.replace("/", "-") if staff else "all_staff"
    return finish_workbook_response(wb, f"ambo_form_e_summary_{safe_id}_{(year or 'all_years').replace('/', '-')}.xlsx")


@login_required
def administrative_create(request):
    if not is_administrative_manager(request.user):
        return render_forbidden(request)
    form = AdministrativeEvaluationForm(request.POST or None)
    form.fields["staff"].queryset = (scoped_staff(request.user) if get_role(request.user) in {"department_head", "committee"} and not request.user.is_superuser else StaffMember.objects.filter(is_active=True)).order_by("name")
    if request.method == "POST" and form.is_valid():
        record = form.save(commit=False)
        record.evaluator = request.user
        record.status = "submitted"
        record.save()
        messages.success(request, "የ6 ወር የአፈፃፀም ምዘናው ተመዝግቧል።")
        return redirect("administrative_detail", pk=record.pk)
    return render(request, "administrative_form.html", {"form": form, "is_edit": False, "staff_directory_data": staff_directory_data(form.fields["staff"].queryset)})


@login_required
def administrative_list(request):
    records = visible_administrative_evaluations(request.user)
    if query := request.GET.get("q", "").strip():
        records = records.filter(Q(staff__name__icontains=query) | Q(staff__employee_id__icontains=query) | Q(staff__department__icontains=query))
    status = request.GET.get("status", "").strip()
    if status in {"draft", "submitted"}:
        records = records.filter(status=status)
    return render(request, "administrative_list.html", {"records": records, "query": request.GET.get("q", ""), "status_filter": status, "can_export": is_manager(request.user), "can_create": is_administrative_manager(request.user)})


@login_required
def administrative_detail(request, pk):
    record = get_object_or_404(visible_administrative_evaluations(request.user), pk=pk)
    return render(request, "administrative_detail.html", {"record": record, "can_archive": can_archive(request.user)})


@login_required
def administrative_update(request, pk):
    record = get_object_or_404(visible_administrative_evaluations(request.user), pk=pk)
    if request.method == "POST":
        form = AdministrativeEvaluationForm(request.POST, instance=record)
        if form.is_valid():
            form.save()
            messages.success(request, "Six-month evaluation updated successfully.")
            return redirect("administrative_detail", pk=record.pk)
    else:
        form = AdministrativeEvaluationForm(instance=record)
    return render(request, "administrative_form.html", {"form": form, "is_edit": True})


@login_required
def staff_list(request):
    if is_manager(request.user):
        staff = scoped_staff(request.user, include_inactive=True)
    else:
        staff = StaffMember.objects.filter(linked_user=request.user)
    if query := request.GET.get("q", "").strip():
        staff = staff.filter(Q(name__icontains=query) | Q(employee_id__icontains=query) | Q(department__icontains=query))
    return render(request, "staff_list.html", {"staff_members": staff, "query": request.GET.get("q", ""), "can_manage": request.user.is_superuser, "can_import": request.user.is_superuser})


@login_required
def staff_create(request):
    if not request.user.is_superuser:
        return render_forbidden(request)
    form = StaffMemberForm(request.POST or None)
    if request.method == "POST" and form.is_valid():
        form.save()
        messages.success(request, "Staff member saved.")
        return redirect("staff_list")
    return render(request, "staff_form.html", {"form": form})


@login_required
@require_POST
def staff_import(request):
    if not request.user.is_superuser:
        return render_forbidden(request)
    uploaded = request.FILES.get("workbook")
    if not uploaded:
        messages.error(request, "Choose an .xlsx file first.")
        return redirect("staff_list")
    if not uploaded.name.lower().endswith(".xlsx"):
        messages.error(request, "Only Excel .xlsx files are supported.")
        return redirect("staff_list")
    try:
        wb = load_workbook(uploaded, read_only=True, data_only=True)
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        headers = [str(v or "").strip().lower().replace(" ", "_") for v in next(rows, [])]
        required = {"employee_id", "name"}
        if not required.issubset(set(headers)):
            messages.error(request, "Excel header row must include employee_id and name. Optional headers: email, college, department, position, grade, job_process.")
            return redirect("staff_list")
        index = {name: i for i, name in enumerate(headers)}
        created, updated, skipped = 0, 0, 0
        for row in rows:
            employee_id = str(row[index["employee_id"]] or "").strip() if index["employee_id"] < len(row) else ""
            name = str(row[index["name"]] or "").strip() if index["name"] < len(row) else ""
            if not employee_id or not name:
                skipped += 1
                continue
            data = {}
            for field in ["email", "college", "department", "position", "grade", "job_process"]:
                if field in index and index[field] < len(row):
                    data[field] = str(row[index[field]] or "").strip()
            if not data.get("college"):
                data["college"] = "Ambo University"
            obj, was_created = StaffMember.objects.update_or_create(employee_id=employee_id, defaults={"name": name, **data})
            if was_created:
                created += 1
            else:
                updated += 1
        messages.success(request, f"Excel import completed: {created} created, {updated} updated, {skipped} blank/invalid row(s) skipped.")
    except Exception as exc:
        messages.error(request, f"Could not read this workbook: {exc}")
    return redirect("staff_list")


# --- Institutional archiving and oversight ------------------------------------
@login_required
def archive_list(request):
    if not is_manager(request.user):
        return render_forbidden(request)
    entries = visible_archive_entries(request.user)
    source_type = request.GET.get("source_type", "").strip()
    level = request.GET.get("level", "").strip()
    query = request.GET.get("q", "").strip()
    if source_type in {"academic", "administrative"}:
        entries = entries.filter(source_type=source_type)
    if level in {"department", "college", "university"}:
        entries = entries.filter(archive_level=level)
    if query:
        entries = entries.filter(
            Q(reference_code__icontains=query)
            | Q(forwarded_to_office__icontains=query)
            | Q(academic_record__staff__name__icontains=query)
            | Q(academic_record__staff__employee_id__icontains=query)
            | Q(administrative_record__staff__name__icontains=query)
            | Q(administrative_record__staff__employee_id__icontains=query)
        )
    return render(request, "archive_list.html", {
        "entries": entries.order_by("-archived_at")[:500],
        "source_type_filter": source_type,
        "level_filter": level,
        "query": query,
        "can_archive": can_archive(request.user),
        "can_export": is_manager(request.user),
    })


@login_required
def archive_create(request, source_type, pk):
    if not can_archive(request.user):
        return render_forbidden(request)
    if source_type == "academic":
        source_record = get_object_or_404(visible_evaluations(request.user), pk=pk)
        staff = source_record.staff
        source_title = f"{FORM_TITLES[source_record.form_type]} — {staff.name}"
        record_url_name = "evaluation_detail"
    elif source_type == "administrative":
        source_record = get_object_or_404(visible_administrative_evaluations(request.user), pk=pk)
        staff = source_record.staff
        source_title = f"የ6 ወር ምዘና — {staff.name}"
        record_url_name = "administrative_detail"
    else:
        return HttpResponseBadRequest("Unknown archive source type.")

    form = ArchiveEntryForm(request.POST or None)
    if request.method == "POST" and form.is_valid():
        entry = form.save(commit=False)
        entry.source_type = source_type
        entry.archived_by = request.user
        if source_type == "academic":
            entry.academic_record = source_record
            entry.administrative_record = None
        else:
            entry.administrative_record = source_record
            entry.academic_record = None
        try:
            entry.full_clean()
            entry.save()
        except ValidationError as exc:
            for error in exc.messages:
                form.add_error(None, error)
        else:
            messages.success(request, "Archive register entry saved. Remember to follow the institution's physical/electronic record-transfer procedure.")
            return redirect("archive_list")
    return render(request, "archive_form.html", {
        "form": form,
        "source_type": source_type,
        "source_record": source_record,
        "staff": staff,
        "source_title": source_title,
        "record_url_name": record_url_name,
        "is_archive_create": True,
    })


@login_required
def export_archive(request):
    if not is_manager(request.user):
        return render_forbidden(request)
    entries = visible_archive_entries(request.user)
    source_type = request.GET.get("source_type", "").strip()
    level = request.GET.get("level", "").strip()
    query = request.GET.get("q", "").strip()
    if source_type in {"academic", "administrative"}:
        entries = entries.filter(source_type=source_type)
    if level in {"department", "college", "university"}:
        entries = entries.filter(archive_level=level)
    if query:
        entries = entries.filter(
            Q(reference_code__icontains=query)
            | Q(forwarded_to_office__icontains=query)
            | Q(academic_record__staff__name__icontains=query)
            | Q(academic_record__staff__employee_id__icontains=query)
            | Q(administrative_record__staff__name__icontains=query)
            | Q(administrative_record__staff__employee_id__icontains=query)
        )
    headers = ["Archive date", "Record type", "Archive level", "Reference code", "Employee ID", "Staff name", "Department", "Period / academic year", "Forwarded to", "Instructor copy sent", "Department copy sent", "Archived by", "Notes"]
    wb, ws = styled_workbook("Archive register", headers)
    for entry in entries.order_by("-archived_at"):
        staff = entry.staff_member
        if entry.source_type == "academic":
            record = entry.academic_record
            period = record.academic_year or f"{record.period_start or ''} – {record.period_end or ''}"
            department = record.department
        else:
            record = entry.administrative_record
            period = f"{record.period_start} – {record.period_end}"
            department = staff.department
        ws.append([
            timezone.localtime(entry.archived_at).strftime("%Y-%m-%d %H:%M"),
            entry.get_source_type_display(), entry.get_archive_level_display(), entry.reference_code,
            staff.employee_id if staff else "", staff.name if staff else "", department, period,
            entry.forwarded_to_office, "Yes" if entry.instructor_copy_sent else "No",
            "Yes" if entry.department_copy_sent else "No",
            entry.archived_by.get_full_name() or entry.archived_by.username, entry.notes,
        ])
    format_sheet(ws)
    return finish_workbook_response(wb, f"ambo_archive_register_{timezone.localdate().isoformat()}.xlsx")
